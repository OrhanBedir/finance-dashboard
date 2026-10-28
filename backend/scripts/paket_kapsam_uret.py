#!/usr/bin/env python3
# BoQ paket / kalem SOW bilgi tabanı üreticisi (28.10.2026, Orhan).
# Kaynak (yalnız okunur): Orhan'ın Mac'indeki BoQ Excel'leri — fiyat ALINMAZ (fiyat boq_katalog'dan, imzalı eklerden gelir).
#   BOQ2026: Türkiye 2026-2027 Integrated Frame BOQ for TE Tender.xlsx → "L4-Package Specification" (paket) + "L4-Item Specification" (kalem)
#            Türkiye 2026-2027 Integrated Frame BOQ for New Site Tender.xlsx → "L4-Specification" (New Site paketleri)
#   BOQ2025: Şimşek/BoQ/BoQ .xlsx → "L4-All Items SOW" (Business Type = Package Item → paket, diğerleri kalem)
# Çıktı: backend/data/boq_paket_kapsam.json — her kayıt: sürüm, tür (PAKET/KALEM), kod, başlık EN/TR, birim, M/S,
# SOW satırları (sıra no + hücre referansı), kabul kriteri, kaynak dosya/sayfa/satır/hücre.
# Etiketleme (dahil/hariç, anten revizyonu, aynı ziyaret …) çalışma anında backend/paketKapsam.js'de yapılır — tek kaynak.
import json, os, re, sys, openpyxl
from openpyxl.utils import get_column_letter as L
B = os.path.expanduser("~/Desktop/Şimşek/BoQ/")
CIKTI = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "boq_paket_kapsam.json")

def satirlar(sow):
    out = []
    for parca in re.split(r"\n|(?:^|\s)\*(?=\S|\s)", str(sow or "")):
        t = parca.strip().lstrip("*").strip()
        if t: out.append(t)
    return out

kayitlar = []
def ekle(surum, tur, kod, en, tr, birim, ms, sow_en, sow_tr, kabul, dosya, sayfa, satir, kol_kod, kol_baslik, kol_sow, is_tipi=None):
    if kod is None or str(kod).strip() in ("", "None"): return
    kayitlar.append({
        "surum": surum, "tur": tur, "kod": str(kod).strip(), "baslik_en": (en or "").strip(), "baslik_tr": (tr or "").strip(),
        "birim": (birim or "").strip(), "ms": (ms or "").strip(), "is_tipi": (is_tipi or "").strip(),
        "sow_satirlari": [{"no": i + 1, "metin": m} for i, m in enumerate(satirlar(sow_en))],
        "sow_tr": (sow_tr or "").strip(), "kabul": (str(kabul).strip() if kabul else ""),
        "kaynak": {"dosya": dosya, "sayfa": sayfa, "satir": satir, "kod_hucre": f"{L(kol_kod)}{satir}", "baslik_hucre": f"{L(kol_baslik)}{satir}", "sow_hucre": f"{L(kol_sow)}{satir}"},
    })

TE = "2026/Türkiye 2026-2027 Integrated Frame BOQ for TE Tender.xlsx"
wb = openpyxl.load_workbook(B + TE, read_only=True, data_only=True)
for sayfa, tur in (("L4-Package Specification", "PAKET"), ("L4-Item Specification", "KALEM")):
    for i, r in enumerate(wb[sayfa].iter_rows(min_row=2, values_only=True), 2):
        if r[1] is None or not re.match(r"^[0-9A-Z]{6,}$", str(r[1]).strip()): continue
        ekle("BOQ2026", tur, r[1], r[3], r[4], r[5], r[6], r[10], r[11], r[14], "Şimşek/BoQ/" + TE, sayfa, i, 2, 4, 11, r[0])
NS = "2026/Türkiye 2026-2027 Integrated Frame BOQ for New Site Tender.xlsx"
for i, r in enumerate(openpyxl.load_workbook(B + NS, read_only=True, data_only=True)["L4-Specification"].iter_rows(min_row=2, values_only=True), 2):
    if r[1] is None or not re.match(r"^[0-9A-Z]{6,}$", str(r[1]).strip()): continue
    ekle("BOQ2026", "PAKET", r[1], r[3], r[4], r[5], r[6], r[10], r[11], None, "Şimşek/BoQ/" + NS, "L4-Specification", i, 2, 4, 11, "New Site")
for i, r in enumerate(openpyxl.load_workbook(B + "BoQ .xlsx", read_only=True, data_only=True)["L4-All Items SOW"].iter_rows(min_row=2, values_only=True), 2):
    if r[3] is None: continue
    tur = "PAKET" if "package" in str(r[2] or "").lower() else "KALEM"
    ekle("BOQ2025", tur, r[3], r[4], None, r[5], r[6], r[7], r[8], r[9], "Şimşek/BoQ/BoQ .xlsx", "L4-All Items SOW", i, 4, 5, 8, r[2])
json.dump({"uretim": "scripts/paket_kapsam_uret.py", "not": "Fiyat içermez; fiyat boq_katalog (imzalı ekler) kaynaklıdır.", "kayitlar": kayitlar},
          open(CIKTI, "w"), ensure_ascii=False, indent=1)
from collections import Counter
print(len(kayitlar), Counter((k["surum"], k["tur"]) for k in kayitlar))
